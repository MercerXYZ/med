"""Экспорт КР: модель документа по приказу 103н, Markdown, Word, полный JSON и JSON по схеме kr_recommendation."""
from __future__ import annotations

import json
import re
import time
from pathlib import Path
from typing import IO

from . import db, prompts as P
from .config import ROOT, SCHEMA_VERSION, SETTINGS
from .groups import group_statement
from .projects import NotFound
from .resolve import excluded_fact_ids, sides
from .synth import intervention_type_for, kr_evidence, udd_for

AGE_RU = {"adults": "Взрослые", "children": "Дети", "mixed": "Взрослые, дети"}
DESIGN_RU = {
    "meta_analysis_rct": "метаанализ РКИ", "systematic_review_rct": "систематический обзор РКИ",
    "rct": "РКИ", "meta_analysis_diagnostic": "метаанализ диагностических исследований",
    "diagnostic_with_reference": "диагностическое исследование с референсным методом",
    "diagnostic_without_reference": "диагностическое исследование без референсного метода",
    "cohort": "когортное исследование", "case_control": "исследование «случай-контроль»",
    "non_comparative": "несравнительное исследование", "case_series": "серия случаев",
    "case_report": "описание случая", "expert_opinion": "мнение экспертов",
    "systematic_review_other": "систематический обзор", "clinical_guideline": "клиническое руководство",
    "narrative_review": "обзор литературы", "textbook": "учебник", "unknown": "дизайн не определён",
}
CHAPTERS = [
    ("1", "Краткая информация по заболеванию или состоянию (группе заболеваний или состояний)",
     ["1.1", "1.2", "1.3", "1.4", "1.5", "1.6"]),
    ("2", "Диагностика заболевания или состояния (группы заболеваний или состояний), медицинские "
          "показания и противопоказания к применению методов диагностики",
     ["2.1", "2.2", "2.3", "2.4", "2.5"]),
    ("3", "Лечение, включая медикаментозную и немедикаментозную терапии, диетотерапию, обезболивание, "
          "медицинские показания и противопоказания к применению методов лечения", ["3.1", "3.2", "3.3"]),
    ("4", "Медицинская реабилитация и санаторно-курортное лечение, медицинские показания и "
          "противопоказания к применению методов медицинской реабилитации", ["4"]),
    ("5", "Профилактика и диспансерное наблюдение, медицинские показания и противопоказания к "
          "применению методов профилактики", ["5"]),
    ("6", "Организация оказания медицинской помощи", ["6"]),
    ("7", "Дополнительная информация (в том числе факторы, влияющие на исход заболевания или состояния)", ["7"]),
]
EMPTY = "Нет данных в загруженных источниках."
NORMATIVE_BASIS = "Приказ Минздрава России от 28.02.2019 № 103н (ред. от 28.09.2023)"
CRITERIA_TITLE = "Критерии оценки качества медицинской помощи"
REFERENCES_TITLE = "Список литературы"
A2_TITLE = "Приложение А2. Методология разработки клинических рекомендаций"
METHODOLOGY = [
    ("Шкала оценки уровней достоверности доказательств (УДД) для методов диагностики", "УДД", [
        ("1", "Систематические обзоры исследований с контролем референсным методом или систематический обзор "
              "рандомизированных клинических исследований с применением метаанализа"),
        ("2", "Отдельные исследования с контролем референсным методом или отдельные рандомизированные клинические "
              "исследования и систематические обзоры исследований любого дизайна, за исключением "
              "рандомизированных клинических исследований, с применением метаанализа"),
        ("3", "Исследования без последовательного контроля референсным методом или исследования с референсным "
              "методом, не являющимся независимым от исследуемого метода, или нерандомизированные "
              "сравнительные исследования, в том числе когортные исследования"),
        ("4", "Несравнительные исследования, описание клинического случая"),
        ("5", "Имеется лишь обоснование механизма действия или мнение экспертов")]),
    ("Шкала оценки уровней достоверности доказательств (УДД) для методов профилактики, лечения и "
     "реабилитации (профилактических, лечебных, реабилитационных вмешательств)", "УДД", [
        ("1", "Систематический обзор рандомизированных клинических исследований с применением метаанализа"),
        ("2", "Отдельные рандомизированные клинические исследования и систематические обзоры исследований "
              "любого дизайна, за исключением рандомизированных клинических исследований, с применением метаанализа"),
        ("3", "Нерандомизированные сравнительные исследования, в т.ч. когортные исследования"),
        ("4", "Несравнительные исследования, описание клинического случая или серии случаев, исследования "
              "«случай-контроль»"),
        ("5", "Имеется лишь обоснование механизма действия вмешательства (доклинические исследования) "
              "или мнение экспертов")]),
    ("Шкала оценки уровней убедительности рекомендаций (УУР) для методов профилактики, диагностики, "
     "лечения и реабилитации (профилактических, диагностических, лечебных, реабилитационных вмешательств)", "УУР", [
        ("A", "Сильная рекомендация (все рассматриваемые критерии эффективности (исходы) являются важными, "
              "все исследования имеют высокое или удовлетворительное методологическое качество, их выводы по "
              "интересующим исходам являются согласованными)"),
        ("B", "Условная рекомендация (не все рассматриваемые критерии эффективности (исходы) являются важными, "
              "не все исследования имеют высокое или удовлетворительное методологическое качество и/или их "
              "выводы по интересующим исходам не являются согласованными)"),
        ("C", "Слабая рекомендация (отсутствие доказательств надлежащего качества (все рассматриваемые "
              "критерии эффективности (исходы) являются неважными, все исследования имеют низкое "
              "методологическое качество и их выводы по интересующим исходам не являются согласованными))")]),
]


# ---------------------------------------------------------------- общее ---

def _project(project_id: str) -> dict:
    p = db.q1("SELECT * FROM projects WHERE id=?", [project_id])
    if not p:
        raise NotFound(project_id)
    return p


def ids_of(v) -> list[str]:
    if isinstance(v, str):
        try:
            v = json.loads(v)
        except json.JSONDecodeError:
            return []
    return [i for i in v or [] if isinstance(i, str)]


def fact_id(f: dict) -> str:
    return f.get("_id") or f.get("id")


def fact_status(row: dict) -> str:
    return row.get("user_status") or row.get("status") or "unverified"


def fact_groups(project_id: str) -> dict[str, str]:
    out = {}
    for g in db.q("SELECT id, card_ids FROM clusters WHERE project_id=?", [project_id]):
        for i in ids_of(g["card_ids"]):
            out.setdefault(i, g["id"])
    return out


def _int(v) -> int | None:
    try:
        return int(v)
    except (TypeError, ValueError):
        return None


def _levels(udd, uur) -> tuple[int, str]:
    # без УДД остаётся только мнение экспертов: УДД 5, УУР C (прил. А2)
    udd = _int(udd)
    return (udd if udd in (1, 2, 3, 4, 5) else 5), (uur if uur in ("A", "B", "C") else "C")


def _grade(uur, udd) -> str:
    return f"Уровень убедительности рекомендаций {uur} (уровень достоверности доказательств – {udd})"


def _ref_text(doc: dict) -> str:
    b = doc.get("biblio") or {}
    authors = [a for a in b.get("authors") or [] if isinstance(a, str) and a.strip()]
    s = ", ".join(authors[:3]) + (" et al." if len(authors) > 3 else "")
    s = f"{s.rstrip('.')}. " if s else ""
    s += (b.get("title") or doc.get("filename") or "").strip()
    if b.get("journal"):
        s += f" // {b['journal']}"
    if b.get("year"):
        s += f". {b['year']}"
    if b.get("doi"):
        s += f". DOI: {b['doi']}"
    return s


def _chapter_key(section: str | None) -> str:
    s = str(section or "")
    if s in P.SECTIONS or s in {c[0] for c in CHAPTERS}:
        return s
    head = s.split(".")[0]
    return head if head in {c[0] for c in CHAPTERS} else "7"


def with_refs(item: dict) -> str:
    """Текст тезиса со ссылками в конце: «Рекомендуется … [1, 2].»"""
    text = (item.get("text") or "").strip()
    if item["type"] != "thesis" or not item.get("refs"):
        return text
    return f"{text.rstrip('.')} [{', '.join(map(str, item['refs']))}]."


# ------------------------------------------------------------ модель КР ---

def _kr_refs(row: dict, facts: dict, include_rejected: bool = False) -> list[str] | None:
    """Факты тезиса, вошедшие в КР; None — тезис в КР не идёт."""
    if (row["status"] == "rejected" and not include_rejected) or row.get("merged_into"):
        return None
    manual = row.get("provenance") == "manual"
    if row.get("orphan") and row["status"] != "approved" and not manual:
        return None
    ids = [e for e in ids_of(row.get("evidence_ids")) if e in facts]
    if not ids and not manual:
        return None
    return ids


def _narrative(text: str, ids: list[str], cite) -> tuple[str, list[int]]:
    used: list[int] = []

    def repl(m):
        ns = cite(re.findall(r"ev_[0-9A-Za-z]+", m.group(0)))
        used.extend(n for n in ns if n not in used)
        return " [" + ", ".join(map(str, ns)) + "]" if ns else ""

    out = re.sub(r"\s*\[[^\]]*ev_[^\]]*\]", repl, text or "")
    out = re.sub(r"\.\s*(\[\d+(?:, \d+)*\])\.?", r" \1.", out)  # «текст. [1].» → «текст [1].»
    out = re.sub(r"\s+([.,;:])", r"\1", out).strip()
    if not used and ids:  # ссылки из текста убраны правкой — ставим в конец
        ns = cite(ids)
        if ns:
            used = ns
            out = f"{out.rstrip('.')} [{', '.join(map(str, ns))}]."
    return out, sorted(used)


def narrative_markers(project_id: str, thesis_id: str, text: str) -> str:
    """[n] из окна 4 → [ev_…] фактов этого абзаца из источника n."""
    m = kr_model(project_id)
    it = next((i for s in m["sections"] for i in s["items"] if i.get("id") == thesis_id), None)
    if not it:
        return text
    refs = {r["n"]: r["fact_ids"] for r in m["references"]}

    def repl(mm):
        ids = [f for n in re.findall(r"\d+", mm.group(0)) for f in refs.get(int(n), []) if f in it["fact_ids"]]
        return " [" + ", ".join(dict.fromkeys(ids)) + "]" if ids else ""

    return re.sub(r"\s*\[\d+(?:\s*,\s*\d+)*\]", repl, text)


def kr_model(project_id: str, include_rejected: bool = False) -> dict:
    """КР как документ: разделы 103н с тезисами и абзацами, литература, критерии качества."""
    p = _project(project_id)
    facts = {fact_id(f): f for fs in kr_evidence(project_id).values() for f in fs}
    card_doc = {r["id"]: r["doc_id"] for r in db.q("SELECT id, doc_id FROM cards WHERE project_id=?", [project_id])}
    docs = {d["id"]: d for d in db.q("SELECT id, filename, biblio FROM documents WHERE project_id=?", [project_id])}

    by_sec: dict[str, list] = {}
    for r in db.q("SELECT * FROM theses WHERE project_id=? ORDER BY created_at, id", [project_id]):
        ids = _kr_refs(r, facts, include_rejected)
        if ids is None:
            continue
        t = db.effective(r)
        by_sec.setdefault(_chapter_key(r.get("section") or t.get("kr_section")), []).append((r, t, ids))
    for items in by_sec.values():
        items.sort(key=lambda x: x[1].get("kind") != "narrative")  # абзацы раздела — перед тезисами

    lit: dict[str, int] = {}
    lit_facts: dict[str, list[str]] = {}

    def cite(ids: list[str]) -> list[int]:
        nums = set()
        for i in ids:
            d = card_doc.get(i)
            if i not in facts or not d:
                continue
            if d not in lit:
                lit[d], lit_facts[d] = len(lit) + 1, []
            if i not in lit_facts[d]:
                lit_facts[d].append(i)
            nums.add(lit[d])
        return sorted(nums)

    def item(r: dict, t: dict, ids: list[str]) -> dict:
        rejected = r["status"] == "rejected"  # виден только в окне 4; без ссылок, нумерация как в Word
        cit = (lambda _: []) if rejected else cite
        common = dict(id=r["id"], section=r.get("section"), fact_ids=ids, flags=list(r.get("review_flags") or []),
                      edited=bool(r.get("override")), status=r["status"], pending=bool(r.get("pending_auto")),
                      provenance=r.get("provenance"))
        if t.get("kind") == "narrative":
            raw = t.get("body") or t.get("statement") or ""
            text, refs = _narrative(raw, ids, cit)
            body = re.sub(r"\s*\[[^\]]*ev_[^\]]*\]", "", raw)
            body = re.sub(r"\s+([.,;:])", r"\1", body).strip()
            return dict(type="text", text=text, body=body, refs=refs, **common)
        phrase = t.get("strength_phrase") or ""
        body = (t.get("body") or "").strip().rstrip(".")
        qc = t.get("quality_criterion") or {}
        criterion = (qc.get("criterion_text") or "").strip() if qc.get("is_criterion") and not rejected and (
            qc.get("auto_eligible") or r["status"] == "approved") else ""
        udd, uur = _levels(t.get("udd"), t.get("uur"))
        return dict(type="thesis", text=t.get("statement") or f"{phrase} {body}.".strip(), phrase=phrase, body=body,
                    udd=udd, uur=uur, grade=_grade(uur, udd), comment=(t.get("comment") or "").strip(),
                    criterion=criterion or None, refs=cit(ids), **common)

    def empty() -> dict:
        return dict(type="text", id=None, section=None, text=EMPTY, refs=[], fact_ids=[], flags=[], edited=False,
                    status=None, pending=False, provenance=None, empty=True)

    def sec(key: str, title: str, level: int, items: list[dict]) -> dict:
        return dict(key=key, title=title, heading=f"{key}. {title}" if level == 1 else f"{key} {title}",
                    level=level, items=items)

    sections = []
    for num, title, subs in CHAPTERS:
        head = [item(*x) for x in by_sec.get(num, [])]
        if subs == [num]:
            sections.append(sec(num, title, 1, head or [empty()]))
            continue
        children = []
        for s in subs:
            its = [item(*x) for x in by_sec.get(s, [])]
            if its:
                children.append(sec(s, P.SECTIONS[s], 2, its))
        sections.append(sec(num, title, 1, head or ([] if children else [empty()])))
        sections += children

    theses = [it for s in sections for it in s["items"] if it["type"] == "thesis"]
    return dict(
        title=p["title"],
        meta=dict(icd_codes=p.get("icd_codes") or [], age_category=p.get("age_category"),
                  age_label=AGE_RU.get(p.get("age_category"), p.get("age_category") or ""),
                  developer_org=p.get("developer_org") or "", year=int(time.strftime("%Y"))),
        sections=sections,
        references=[dict(n=n, text=_ref_text(docs.get(d) or {}), doc_id=d, fact_ids=lit_facts[d])
                    for d, n in sorted(lit.items(), key=lambda x: x[1])],
        criteria=[dict(text=t["criterion"], thesis_id=t["id"], udd=t["udd"], uur=t["uur"])
                  for t in theses if t["criterion"]],
    )


# ------------------------------------------------------------- Markdown ---

def _md_cell(s: str) -> str:
    return str(s or "").replace("|", "\\|").replace("\n", " ")


def to_markdown(project_id: str, model: dict | None = None) -> str:
    m = model or kr_model(project_id)
    meta = m["meta"]
    L = [f"# {m['title']}", "", "Клинические рекомендации", "",
         f"Кодирование по МКБ-10: {', '.join(meta['icd_codes'])}  ",
         f"Возрастная группа: {meta['age_label']}  ",
         f"Год утверждения: {meta['year']}  "]
    if meta["developer_org"]:
        L.append(f"Разработчик клинической рекомендации: {meta['developer_org']}  ")
    L.append("")
    for s in m["sections"]:
        L += [("## " if s["level"] == 1 else "### ") + s["heading"], ""]
        for it in s["items"]:
            if it["type"] == "text":
                L += [it["text"], ""]
                continue
            L += [f"- {with_refs(it)}  ", f"  **{it['grade']}**  "]
            if it["comment"]:
                L.append(f"  Комментарий: {it['comment']}")
            L.append("")
    L += [f"## {CRITERIA_TITLE}", ""]
    if m["criteria"]:
        L += ["| № | Критерии качества | Оценка выполнения |", "|---|---|---|"]
        L += [f"| {i} | {_md_cell(c['text'])} | Да/Нет |" for i, c in enumerate(m["criteria"], 1)]
    else:
        L.append(EMPTY)
    L += ["", f"## {REFERENCES_TITLE}", ""]
    L += [f"{r['n']}. {r['text']}" for r in m["references"]] or [EMPTY]
    L += ["", f"## {A2_TITLE}", ""]
    for caption, head, rows in METHODOLOGY:
        L += [f"**{caption}**", "", f"| {head} | Расшифровка |", "|---|---|"]
        L += [f"| {a} | {_md_cell(b)} |" for a, b in rows] + [""]
    return "\n".join(L).rstrip() + "\n"


# ----------------------------------------------------------------- Word ---

_XML_BAD = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")


def xml_safe(text) -> str:
    """Управляющие символы из PDF недопустимы в XML документа Word."""
    return _XML_BAD.sub(" ", str(text or ""))


def to_docx(project_id: str, path: Path | IO[bytes], model: dict | None = None) -> Path | IO[bytes]:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    m = model or kr_model(project_id)
    meta = m["meta"]
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    for name in ("Heading 1", "Heading 2"):
        doc.styles[name].font.name = "Times New Roman"
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)

    def para(text="", bold=False, italic=False, size=None, align=None, style=None):
        par = doc.add_paragraph(style=style)
        run = par.add_run(xml_safe(text))
        run.bold, run.italic = bold, italic
        if size:
            run.font.size = Pt(size)
        if align is not None:
            par.alignment = align
        return par

    def table(header: tuple, rows: list[tuple]):
        tbl = doc.add_table(rows=1, cols=len(header))
        tbl.style = "Table Grid"
        for cell, h in zip(tbl.rows[0].cells, header):
            cell.text = h
        for row in rows:
            for cell, v in zip(tbl.add_row().cells, row):
                cell.text = xml_safe(v)

    center = WD_ALIGN_PARAGRAPH.CENTER
    para("Клинические рекомендации", bold=True, size=16, align=center)
    para(m["title"], bold=True, size=16, align=center)
    para()
    para("Кодирование по Международной статистической классификации болезней и проблем, "
         "связанных со здоровьем: " + ", ".join(meta["icd_codes"]))
    para(f"Возрастная группа: {meta['age_label']}")
    para(f"Год утверждения: {meta['year']}")
    if meta["developer_org"]:
        para(f"Разработчик клинической рекомендации: {meta['developer_org']}")
    doc.add_page_break()

    for s in m["sections"]:
        doc.add_heading(xml_safe(s["heading"]), level=s["level"])
        for it in s["items"]:
            if it["type"] == "text":
                para(it["text"])
                continue
            para(with_refs(it), style="List Bullet")
            para(it["grade"], bold=True)
            if it["comment"]:
                cp = doc.add_paragraph()
                cp.add_run("Комментарий: ").italic = True
                cp.add_run(xml_safe(it["comment"])).italic = True

    doc.add_heading(CRITERIA_TITLE, level=1)
    if m["criteria"]:
        table(("№", "Критерии качества", "Оценка выполнения"),
              [(str(i), c["text"], "Да/Нет") for i, c in enumerate(m["criteria"], 1)])
    else:
        para(EMPTY)
    doc.add_heading(REFERENCES_TITLE, level=1)
    for r in m["references"]:
        para(f"{r['n']}. {r['text']}")
    if not m["references"]:
        para(EMPTY)
    doc.add_heading(A2_TITLE, level=1)
    for caption, head, rows in METHODOLOGY:
        para(caption, bold=True)
        table((head, "Расшифровка"), rows)
        para()
    doc.save(path)
    return path


# ------------------------------------------------------------ полный JSON ---

THESIS_FIELDS = ("subsection_title", "statement", "strength_phrase", "body", "direction", "intervention",
                 "evidence_backed_core", "expert_additions", "comment", "udd", "uur", "uur_preliminary",
                 "grade_rationale", "consensus", "quality_criterion", "explicit_source_recommendation",
                 "expert_author")


def build_export(project_id: str, model: dict | None = None) -> dict:
    from .pipeline import processing_report

    p = _project(project_id)
    m = model or kr_model(project_id)
    items = [it for s in m["sections"] for it in s["items"] if it["id"]]
    cited = {i for it in items for i in it["fact_ids"]}
    in_kr = {it["id"] for it in items}
    in_criteria = {c["thesis_id"] for c in m["criteria"]}
    kr = {gid: [fact_id(f) for f in fs] for gid, fs in kr_evidence(project_id).items()}
    excluded = excluded_fact_ids(project_id)
    groups_of = fact_groups(project_id)
    cards = db.q("SELECT * FROM cards WHERE project_id=? ORDER BY doc_id, id", [project_id])
    by_id = {r["id"]: r for r in cards}

    conflicts, cf_of = [], {}
    for c in db.q("SELECT * FROM conflicts WHERE project_id=? ORDER BY created_at, id", [project_id]):
        a, b = sides(c)
        for i in a + b:
            cf_of.setdefault(i, []).append(c["id"])
        first = next((by_id[i] for i in a + b if i in by_id), None)
        conflicts.append(dict(
            id=c["id"], section=c.get("section") or (db.effective(first).get("kr_section") if first else None) or "7",
            kind=c.get("kind") if c.get("kind") in P.CONFLICT_KINDS else "other",
            description=c.get("description") or "", side_a=a, side_b=b, decision=c.get("decision"),
            decided_by=c.get("decided_by"), reason=c.get("reason"), status=c.get("status") or "open",
            origin=c.get("origin") or "auto"))

    facts = []
    for r in cards:
        c = db.effective(r)
        v = r.get("verification") or {}
        q = v.get("quote") or {}
        facts.append(dict(
            id=r["id"], doc_id=r["doc_id"], statement_ru=c.get("statement_ru") or "",
            claim_type=c.get("claim_type"), polarity=c.get("polarity"), negation=c.get("polarity") == "negative",
            population=c.get("population") or None, age_group=c.get("age_group"),
            intervention=c.get("intervention") or None, comparator=c.get("comparator") or None,
            outcome=c.get("outcome") or None, conditions=c.get("conditions") or None, dose=c.get("dose") or None,
            timing=c.get("timing") or None, thresholds=c.get("thresholds") or None,
            intervention_type=c.get("intervention_type"), kr_section=c.get("kr_section"),
            topic_relevance=c.get("topic_relevance"), secondary_citation=bool(c.get("secondary_citation")),
            source_quote=c.get("quote") or "", pdf_page=_int(q.get("found_page") or c.get("page")),
            verification=dict(status=r["status"], user_status=r.get("user_status"), quote_match=q.get("status"),
                              quote_score=q.get("score"), quote_found_page=_int(q.get("found_page")),
                              checks=v.get("checks") or [], judge=v.get("judge"), issues=v.get("issues") or []),
            group_id=groups_of.get(r["id"]), conflict_ids=cf_of.get(r["id"], []),
            excluded_by_conflict=r["id"] in excluded, in_kr=r["id"] in cited,
            origin=r.get("origin") or "auto", edited_by_expert=bool(r.get("override")), auto_value=r.get("auto"),
            pending_auto_update=bool(r.get("pending_auto")), run_id=r.get("run_id")))

    groups = []
    for g in db.q("SELECT * FROM clusters WHERE project_id=? ORDER BY section, id", [project_id]):
        ov = g.get("override") or {}
        groups.append(dict(
            id=g["id"], section=g.get("section") or "7", statement=group_statement(g),
            auto_statement=group_statement({**g, "override": None}),
            edited_by_expert=bool(isinstance(ov, dict) and ov.get("statement")), excluded=bool(g.get("excluded")),
            card_ids=ids_of(g.get("card_ids")), kr_card_ids=kr.get(g["id"], []), origin=g.get("origin") or "auto",
            run_id=g.get("run_id")))

    theses, qc = [], []
    for r in db.q("SELECT * FROM theses WHERE project_id=? ORDER BY section, created_at", [project_id]):
        t = db.effective(r)
        rec = dict(id=r["id"], section=r.get("section"), group_id=r.get("cluster_id"),
                   kind=t.get("kind") or "recommendation", **{k: t.get(k) for k in THESIS_FIELDS})
        rec.update(statement=rec["statement"] or "", evidence_ids=ids_of(r.get("evidence_ids")), status=r["status"],
                   provenance=r.get("provenance") or "auto", review_flags=list(r.get("review_flags") or []),
                   orphan=bool(r.get("orphan")), in_kr=r["id"] in in_kr, edited_by_expert=bool(r.get("override")),
                   auto_value=r.get("auto"), pending_auto_update=bool(r.get("pending_auto")))
        theses.append(rec)
        crit = t.get("quality_criterion") or {}
        if crit.get("is_criterion") and crit.get("criterion_text") and r["status"] != "rejected":
            qc.append(dict(id="qc_" + r["id"][4:], thesis_id=r["id"], criterion_text=crit["criterion_text"],
                           applicable_stage=crit.get("applicable_stage"), auto_eligible=crit.get("auto_eligible"),
                           in_kr=r["id"] in in_criteria))

    docs = []
    for d in db.q("SELECT * FROM documents WHERE project_id=? ORDER BY filename", [project_id]):
        docs.append(dict(id=d["id"], filename=d["filename"], sha256=d["sha256"], n_pages=d["n_pages"] or 0,
                         processing_status=d["status"], status_detail=d["status_detail"],
                         duplicate_of=d["duplicate_of"],
                         bibliography={k: v for k, v in (d.get("biblio") or {}).items() if not k.startswith("_")}))

    return {
        "schema_version": SCHEMA_VERSION,
        "generated_at": time.strftime("%Y-%m-%dT%H:%M:%S"),
        "generator": {"model": SETTINGS.model, "judge_model": SETTINGS.judge_model,
                      "settings": SETTINGS.to_json() | {"cli_path": None}},
        "project": dict(id=p["id"], title=p["title"], age_category=p["age_category"],
                        icd_codes=p.get("icd_codes") or [], developer_org=p.get("developer_org") or "",
                        normative_basis=NORMATIVE_BASIS),
        "documents": docs,
        "evidence_cards": facts,
        "groups": groups,
        "conflicts": conflicts,
        "recommendations": theses,
        "quality_criteria": qc,
        "kr": m,
        "expert_edits": db.q("SELECT entity, entity_id, action, before, after, comment, ts FROM expert_edits "
                             "WHERE project_id=? ORDER BY ts", [project_id]),
        "processing_report": processing_report(project_id),
    }


# --------------------------------------------------- схема kr_recommendation ---

KRS_DESIGNS = {"meta_analysis_rct", "systematic_review_rct", "rct", "meta_analysis_diagnostic",
                   "diagnostic_with_reference", "diagnostic_without_reference", "cohort", "case_control",
                   "non_comparative", "case_series", "case_report", "expert_opinion"}
DESIGN_TO_KRS = {"systematic_review_other": "systematic_review_rct"}  # остальное вне шкалы → expert_opinion
PHRASES = ("Рекомендуется", "Рекомендуется рассмотреть возможность", "Не рекомендуется", "Не рекомендуется рассматривать")
ITYPES = ("diagnostic", "therapeutic", "preventive", "rehab")
STAGE_OF = {"diagnostic": "diagnosis", "therapeutic": "treatment", "rehab": "rehabilitation",
            "preventive": "prevention"}
ADDITION_BASIS = ("expert_consensus", "evidence+expert", "guideline_cross_ref")


def _krs_section(section: str | None) -> str:
    s = str(section or "")
    if s in P.SECTIONS:
        return s
    head = s.split(".")[0]
    return next((k for k in P.SECTIONS if k.split(".")[0] == head), "7")


def _krs_rec(p: dict, row: dict, t: dict, it: dict, refs: list[str]) -> dict:
    section = _krs_section(row.get("section") or t.get("kr_section"))
    iv = t.get("intervention") if isinstance(t.get("intervention"), dict) else {}
    itype = iv.get("type") if iv.get("type") in ITYPES else intervention_type_for(section)
    phrase = it["phrase"] if it["phrase"] in PHRASES else None
    direction = ("negative" if phrase.startswith("Не") else "positive") if phrase else (
        t.get("direction") if t.get("direction") in ("positive", "negative") else "positive")
    why = (t.get("grade_rationale") or "").strip() or "Обоснование не указано"
    if _int(t.get("udd")) not in (1, 2, 3, 4, 5):
        why += "; УДД не определён, указан 5"
    if t.get("uur") not in ("A", "B", "C"):
        why += "; УУР не определён, указан C"
    rec = dict(id=row["id"], kr_section=section, statement=it["text"], direction=direction,
               intervention=dict(type=itype, scale_ref="udd_diagnostic" if itype == "diagnostic" else "udd_treatment",
                                 **{k: v for k, v in iv.items()
                                    if k in ("action", "target_population", "outcome") and isinstance(v, str)},
                                 **{k: v for k, v in iv.items()
                                    if k in ("comparator", "mnn", "atc_code") and (v is None or isinstance(v, str))}),
               evidence_refs=refs, udd=it["udd"], uur=it["uur"], grade_rationale=why,
               consensus={"agreement": "consistent", "conflicting_refs": []},
               icd_codes=p.get("icd_codes") or [],
               provenance=row.get("provenance") if row.get("provenance") in ("auto", "auto+expert_review", "manual")
               else "auto",
               review_flags=list(it["flags"]))
    if phrase:
        rec["strength_phrase"] = phrase
    if isinstance(t.get("subsection_title"), str) and t["subsection_title"]:
        rec["subsection_title"] = t["subsection_title"]
    if isinstance(t.get("evidence_backed_core"), str) and t["evidence_backed_core"]:
        rec["evidence_backed_core"] = t["evidence_backed_core"]
    rec["expert_additions"] = [
        dict(claim=str(a.get("claim") or ""),
             basis=a.get("basis") if a.get("basis") in ADDITION_BASIS else "expert_consensus",
             refs=[r for r in a.get("refs") or [] if r in refs])
        for a in t.get("expert_additions") or [] if isinstance(a, dict) and a.get("claim")]
    qc = t.get("quality_criterion") if isinstance(t.get("quality_criterion"), dict) else {}
    if qc:
        out = dict(is_criterion=bool(qc.get("is_criterion")), criterion_text=qc.get("criterion_text") or None)
        if qc.get("applicable_stage") in STAGE_OF.values():
            out["applicable_stage"] = qc["applicable_stage"]
        for k in ("measurable", "auto_eligible"):
            if isinstance(qc.get(k), bool):
                out[k] = qc[k]
        if out["is_criterion"]:
            out.setdefault("applicable_stage", STAGE_OF[itype])
            out.setdefault("measurable", bool(out["criterion_text"]))
            if not out["criterion_text"]:
                out["is_criterion"] = False
        rec["quality_criterion"] = out
    return rec


def _krs_unit(row: dict, doc: dict) -> dict:
    c = db.effective(row)
    b = doc.get("biblio") or {}
    section = _krs_section(c.get("kr_section") or "7")
    design = b.get("study_design")
    src = dict(title=(b.get("title") or doc.get("filename") or "Без названия").strip(),
               study_design=design if design in KRS_DESIGNS else DESIGN_TO_KRS.get(design, "expert_opinion"))
    for k in ("doi", "journal", "population"):
        if isinstance(b.get(k), str) and b[k].strip():
            src[k] = b[k].strip()
    if str(b.get("pmid") or "").isdigit():
        src["pmid"] = str(b["pmid"])
    authors = [a for a in b.get("authors") or [] if isinstance(a, str) and a.strip()]
    if authors:
        src["authors"] = authors
    year = _int(b.get("year"))
    if year and 1900 <= year <= 2100:
        src["year"] = year
    size = _int(b.get("sample_size"))
    if size and size > 0:
        src["sample_size"] = size
    udd = None if c.get("secondary_citation") else udd_for(design, section)
    v = row.get("verification") or {}
    notes = list(v.get("issues") or [])
    if row.get("user_status") == "confirmed":
        notes.insert(0, "подтверждено экспертом")
    if udd is None:
        notes.append("УДД источника не определён, указан 5")
    statement = (c.get("statement_ru") or "").strip() or "—"
    verified = row.get("user_status") == "confirmed" or (row.get("status") == "verified" and not row.get("user_status"))
    return dict(
        id=row["id"], subject=(c.get("intervention") or "").strip() or statement[:80],
        predicate=c.get("polarity") or "descriptive", object=(c.get("outcome") or "").strip() or statement,
        negation=c.get("polarity") == "negative", source=src, intervention_type=intervention_type_for(section),
        udd_source=udd or 5, verification_status="verified" if verified else "unverified",
        verification_notes="; ".join(notes) or "проверки пройдены", kr_section=section,
        statement_ru=statement, source_quote=c.get("quote") or "",
        pdf_page=_int(((v.get("quote") or {}).get("found_page")) or c.get("page")))


def kr_schema_export(project_id: str, model: dict | None = None) -> tuple[dict, list[str]]:
    """Экспорт по kr_recommendation.schema.json: (документ, ошибки валидации)."""
    p = _project(project_id)
    m = model or kr_model(project_id)
    rows = {r["id"]: r for r in db.q("SELECT * FROM theses WHERE project_id=?", [project_id])}
    recs, used = [], []
    for s in m["sections"]:
        for it in s["items"]:
            if it["type"] != "thesis" or it["id"] not in rows:
                continue
            row = rows[it["id"]]
            refs = list(it["fact_ids"]) or ["ev_expert_" + re.sub(r"[^0-9A-Za-z_]", "_", row["id"][4:])]
            used += [e for e in refs if e not in used]
            recs.append(_krs_rec(p, row, db.effective(row), it, refs))

    cards = {r["id"]: r for r in db.q("SELECT * FROM cards WHERE project_id=?", [project_id])}
    docs = {d["id"]: d for d in db.q("SELECT id, filename, biblio FROM documents WHERE project_id=?", [project_id])}
    units = []
    for eid in used:
        if eid in cards:
            units.append(_krs_unit(cards[eid], docs.get(cards[eid]["doc_id"]) or {}))
        else:
            units.append(dict(id=eid, subject="Экспертное мнение", predicate="обосновывает",
                              object="тезис без опоры на статьи",
                              source=dict(title="Экспертное дополнение", study_design="expert_opinion"),
                              udd_source=5, verification_status="unverified",
                              verification_notes="экспертное дополнение, цитаты нет"))
    out = {"kr_meta": dict(title=p["title"], year=int(time.strftime("%Y")), icd_codes=p.get("icd_codes") or [],
                           age_category=p["age_category"], developer_org=p.get("developer_org") or "",
                           schema_version="1.0"),
           "evidence_units": units, "recommendation_theses": recs}
    return out, _validate(out, ROOT / "kr_recommendation.schema.json")


def _validate(data: dict, schema_path: Path) -> list[str]:
    if not schema_path.exists():
        return []
    import jsonschema
    schema = json.loads(schema_path.read_text(encoding="utf-8"))
    v = jsonschema.Draft202012Validator(schema)
    return [f"{'/'.join(map(str, e.absolute_path))}: {e.message}" for e in v.iter_errors(data)]


def export_all(project_id: str, out_dir: Path) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    base = out_dir / f"{project_id}_{time.strftime('%Y%m%d_%H%M')}"
    m = kr_model(project_id)
    data = json.loads(json.dumps(build_export(project_id, m), ensure_ascii=False, default=str))
    full = Path(f"{base}.json")
    full.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    own = _validate(data, ROOT / "schemas" / "krgen_export.schema.json")
    Path(f"{base}.validation.txt").write_text(
        "\n".join(own[:200]) + "\n" if own else "OK: соответствует schemas/krgen_export.schema.json\n", encoding="utf-8")
    krs, errors = kr_schema_export(project_id, m)
    tp = Path(f"{base}_kr_recommendation.json")
    tp.write_text(json.dumps(krs, ensure_ascii=False, indent=1), encoding="utf-8")
    Path(f"{base}_kr_recommendation.validation.txt").write_text(
        "\n".join(errors) + "\n" if errors else "OK: соответствует kr_recommendation.schema.json\n", encoding="utf-8")
    md = Path(f"{base}.md")
    md.write_text(to_markdown(project_id, m), encoding="utf-8")
    return [full, tp, md, to_docx(project_id, Path(f"{base}.docx"), m)]
